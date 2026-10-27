"""[A] PRD 11장 수용 기준 + 13장 결정사항 검증."""

from __future__ import annotations

import csv
import io
import re
import subprocess
import sys
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from playwright.sync_api import Page

from src.charts import HEATMAP_COLUMNS
from src.components import DEADLINE_ALERT_DAYS, closing_soon, format_d_day
from src.filters import Filters, TARGET_LABELS, apply_filters, target_level
from src.models import (
    CONSULTATION_TYPE_LABELS,
    LEVEL_LABELS,
    REQUIREMENT_LABELS,
    STAGE_LABELS,
    consultation_status,
    days_until_end,
)
from tests.e2e.conftest import (
    BY_ID,
    DATA,
    RESULTS,
    ROOT,
    TODAY,
    goto,
    main_text,
    multiselect,
    norm,
    record,
    settle,
    shot,
    sidebar_text,
)

PAGES = {"dashboard": "", "list": "policies", "detail": "detail", "consult": "consultations",
         "impact": "impact", "changelog": "changelog"}
N_IN_FORCE = sum(p.status == "in_force" for p in DATA.policies)
N_PENDING = sum(p.status == "pending" for p in DATA.policies)


# A1 -------------------------------------------------------------------------
def test_a1_list_split_into_two_tabs(page: Page, app_url):
    goto(page, app_url, "policies")
    shot(page, "A1_list_tabs")
    assert page.get_by_role("tab").all_inner_texts() == [
        f"시행 중인 법·정책 ({N_IN_FORCE})",
        f"제정 중·시행 예정 ({N_PENDING})",
    ]
    in_force_titles = {norm(t) for t in page.locator("[role=tabpanel]:visible .card-title").all_inner_texts()}
    assert in_force_titles == {p.name_ko for p in DATA.policies if p.status == "in_force"}
    page.get_by_role("tab", name=re.compile("제정 중")).click()
    settle(page)
    pending_titles = {norm(t) for t in page.locator("[role=tabpanel]:visible .card-title").all_inner_texts()}
    assert pending_titles == {p.name_ko for p in DATA.policies if p.status == "pending"}


# A2·A3·A5·A6 (정책별 상세) ------------------------------------------------------
@pytest.mark.parametrize("pid", list(BY_ID))
def test_a2_a3_policy_detail(page: Page, app_url, pid):
    p = BY_ID[pid]
    goto(page, app_url, f"detail?policy={pid}")
    shot(page, f"A2_detail_{pid}")
    text = main_text(page)
    problems = []

    if page.locator("[data-testid=stMain] h2").first.inner_text().strip() != p.name_ko:
        problems.append("헤더 정책명 불일치")

    # A2 요구사항 + 유형 태그
    for r in p.requirements:
        if REQUIREMENT_LABELS[r.type] not in text:
            problems.append(f"요구사항 태그 누락: {REQUIREMENT_LABELS[r.type]}")
        if norm(r.text) not in text:
            problems.append(f"요구사항 본문 누락: {r.text[:30]}")

    # A3 제정 중 법안: 입법 단계 스텝, 일정, 의견수렴
    stepper = page.locator("ol.stepper li.current")
    if p.status == "pending":
        if stepper.count() != 1:
            problems.append("입법 단계 스텝 없음")
        elif STAGE_LABELS[p.stage] not in norm(stepper.inner_text()):
            problems.append(f"현재 단계 불일치: {stepper.inner_text()!r} vs {STAGE_LABELS[p.stage]}")
    elif stepper.count():
        problems.append("시행 중 정책에 입법 단계 스텝이 표시됨")
    for m in p.milestones:
        if f"{m.date.isoformat()} — {norm(m.label)}" not in text:
            problems.append(f"마일스톤 누락: {m.date} {m.label[:20]}")
    for c in p.consultations:
        if norm(c.title) not in text:
            problems.append(f"의견수렴 제목 누락: {c.title[:30]}")
        if f"기간 {c.start.isoformat()} ~ {c.end.isoformat()}" not in text:
            problems.append(f"의견수렴 기간 누락: {c.start}~{c.end}")
        if CONSULTATION_TYPE_LABELS[c.type] not in text:
            problems.append(f"의견수렴 유형 누락: {c.type}")
    hys_links = page.locator("[data-testid=stMain] a", has_text="Have Your Say").evaluate_all("els => els.map(e => e.href)")
    for c in p.consultations:
        if c.url and c.url not in hys_links:
            problems.append(f"의견수렴 링크 누락: {c.url}")

    # A5 역내·역외·한국 영향 비교
    for target, level in (("eu", p.impact.eu_producers.level), ("non_eu", p.impact.non_eu_producers.level),
                          ("korea", p.impact.korea_effective_level)):
        if TARGET_LABELS[target] not in text:
            problems.append(f"영향 열 누락: {TARGET_LABELS[target]}")
    levels_shown = re.findall(r"영향도 ([상중하])", text)
    expected_levels = [LEVEL_LABELS[x] for x in (p.impact.eu_producers.level, p.impact.non_eu_producers.level,
                                                 p.impact.korea_effective_level)]
    if levels_shown != expected_levels:
        problems.append(f"영향도 불일치: {levels_shown} vs {expected_levels}")

    # A6 공식 링크 (Commission 필수)
    hrefs = page.locator("[data-testid=stMain] a[target=_blank]").evaluate_all("els => els.map(e => e.href)")
    for field, url in p.links.model_dump().items():
        if url and url not in hrefs:
            problems.append(f"{field} 링크 누락: {url}")

    # D 정합성: 확인 필요 배지, 마지막 확인일
    badges = page.locator("[data-testid=stMain] .stMarkdownBadge").all_inner_texts()
    has_verify = any("확인 필요" in b and "재확인" not in b for b in badges)
    if has_verify != p.verify_needed:
        problems.append(f"확인 필요 배지 표시 오류 (verify_needed={p.verify_needed})")
    if f"마지막 확인일 {p.last_verified.isoformat()}" not in text:
        problems.append("마지막 확인일 누락")

    assert not problems, problems


# A4 -------------------------------------------------------------------------
def test_a4_closing_consultations_in_sidebar_with_d_day(page: Page, app_url):
    expected = closing_soon(DATA.policies, TODAY)
    goto(page, app_url, "")
    shot(page, "A4_sidebar_alerts")
    items = [norm(t) for t in page.locator("[data-testid=stSidebar] .alert-item").all_inner_texts()]
    record("A4_sidebar_alerts", {"expected": [(p.id, c.title, d) for p, c, d in expected], "shown": items})
    assert len(items) == len(expected)
    for (p, c, d), shown in zip(expected, items):
        assert shown.startswith(f"{format_d_day(d)} {p.name_ko}"), shown
        assert CONSULTATION_TYPE_LABELS[c.type] in shown and c.end.isoformat() in shown
    # 의견수렴 페이지의 D-day도 오늘 기준으로 일치
    goto(page, app_url, "consultations")
    boxes = [norm(t) for t in page.locator("[role=tabpanel]:visible .dday").all_inner_texts()]
    open_items = sorted(
        [(c, days_until_end(c, TODAY)) for p in DATA.policies for c in p.consultations
         if consultation_status(c, TODAY) == "open"], key=lambda x: x[0].end)
    assert [b.split()[0] for b in boxes] == [format_d_day(d) for _, d in open_items]


def test_a4_no_alert_outside_window(page: Page, custom_server):
    """마감 31일 남은 의견수렴은 알림에 나오지 않고, 30일·0일 남은 것은 나온다."""
    def mutate(policies):
        cons = policies[0].setdefault("consultations", [])
        for days in (0, DEADLINE_ALERT_DAYS, DEADLINE_ALERT_DAYS + 1):
            cons.append({"type": "public_consultation", "title": f"경계 테스트 {days}일",
                         "start": TODAY - timedelta(days=5), "end": TODAY + timedelta(days=days)})
        for p in policies[1:]:
            p["consultations"] = [c for c in p.get("consultations", []) if c["end"] < TODAY]
        return policies

    url = custom_server(mutate)
    goto(page, url, "")
    shot(page, "A4_alert_window_boundary")
    items = [norm(t) for t in page.locator("[data-testid=stSidebar] .alert-item").all_inner_texts()]
    assert [i.split()[0] for i in items] == ["D-Day", f"D-{DEADLINE_ALERT_DAYS}"]


# A5 -------------------------------------------------------------------------
def test_a5_impact_matrix_matches_yaml(page: Page, app_url):
    goto(page, app_url, "impact")
    shot(page, "A5_impact_matrix")
    heat = page.locator(".js-plotly-plot").first.evaluate(
        "gd => ({x: gd.data[0].x, names: gd.data[0].customdata.map(r => r[0]), text: gd.data[0].text})")
    policies = apply_filters(DATA.policies, Filters())
    assert heat["x"] == list(HEATMAP_COLUMNS.values())
    assert heat["names"] == [p.name_ko for p in policies]
    assert heat["text"] == [[LEVEL_LABELS[target_level(p, t)] for t in TARGET_LABELS] for p in policies]


# A6 -------------------------------------------------------------------------
@pytest.mark.network
def test_a6_no_broken_links():
    proc = subprocess.run([sys.executable, "scripts/check_links.py"], cwd=ROOT, capture_output=True,
                          text=True, encoding="utf-8", timeout=600)
    (RESULTS / "A6_check_links.txt").write_text(proc.stdout + proc.stderr, encoding="utf-8")
    summary = [line for line in proc.stdout.splitlines() if line.startswith("정상")]
    assert proc.returncode == 0, summary


# A7 -------------------------------------------------------------------------
def test_a7_stale_badge_boundary(page: Page, custom_server):
    """last_verified가 13·14일 전이면 배지 없음, 15일 전이면 '재확인 필요'."""
    targets = [p.id for p in DATA.policies if p.status == "in_force"][:3]
    ages = dict(zip(targets, (13, 14, 15)))

    def mutate(policies):
        for p in policies:
            p["last_verified"] = TODAY - timedelta(days=ages.get(p["id"], 0))
        return policies

    url = custom_server(mutate)
    goto(page, url, "policies")
    shot(page, "A7_stale_boundary")
    for pid, age in ages.items():
        card = page.locator(f".st-key-card_{pid}")
        assert ("재확인 필요" in card.inner_text()) is (age > 14), f"{pid}: {age}일"
    # help 툴팁이 있는 배지는 DOM에 2번 렌더링되므로 배지 요소가 아니라 '배지가 있는 카드' 수를 센다
    stale_cards = page.locator("[role=tabpanel]:visible [class*='st-key-card_']").filter(has_text="재확인 필요").count()
    assert stale_cards == 1


# A8 -------------------------------------------------------------------------
def _download(page: Page, label: str):
    # help 툴팁이 있는 버튼은 DOM에 button이 2개 렌더링되므로 접근성 이름(아이콘 'download' + 라벨)으로 찾는다
    with page.expect_download() as info:
        page.get_by_role("button", name=re.compile(rf"{label}$")).first.click()
    return info.value


def _shown_count(page: Page) -> int:
    return sum(int(re.search(r"\((\d+)\)", t).group(1)) for t in page.get_by_role("tab").all_inner_texts())


@pytest.mark.parametrize("relevance", [None, "상"])
def test_a8_export_matches_screen(page: Page, app_url, relevance, tmp_path):
    goto(page, app_url, "policies")
    if relevance:
        multiselect(page, "철강 관련도", [relevance])
    shown = _shown_count(page)

    csv_file = _download(page, "CSV")
    raw = (tmp_path / "out.csv")
    csv_file.save_as(raw)
    data = raw.read_bytes()
    assert data.startswith(b"\xef\xbb\xbf"), "CSV에 UTF-8 BOM이 없어 Excel에서 한글이 깨질 수 있음"
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == shown
    assert all(re.search(r"[가-힣]", r["정책명"]) for r in rows)

    xlsx_file = _download(page, "Excel")
    xlsx = tmp_path / "out.xlsx"
    xlsx_file.save_as(xlsx)
    ws = load_workbook(xlsx).active
    assert ws.max_row - 1 == shown
    assert [c.value for c in ws[1]][:2] == ["ID", "정책명"]
    record(f"A8_export_{relevance or 'all'}", {"shown": shown, "csv_rows": len(rows), "xlsx_rows": ws.max_row - 1,
                                               "csv_file": csv_file.suggested_filename})


# A9·A10 ---------------------------------------------------------------------
@pytest.mark.parametrize("name,path", PAGES.items())
def test_a9_a10_internal_disclaimer_and_no_calculator(page: Page, app_url, name, path):
    goto(page, app_url, path)
    assert "사내용" in sidebar_text(page)
    assert page.locator("[data-testid=stAppDeployButton]").count() == 0
    text = main_text(page)
    assert page.locator("[data-testid=stNumberInput]").count() == 0
    for word in ("계산기", "비용 계산", "탄소발자국 계산", "인증서 비용 산정"):
        assert word not in text, word
